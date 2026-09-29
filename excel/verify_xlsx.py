"""엑셀 파일의 수식이 웹앱 엔진(app/engine.js)과 똑같이 판정하는지 검증한다. 엑셀·LibreOffice 없이 동작.

사용: python excel/verify_xlsx.py ["OX chess_v3.xlsx"]
필요: node, 파이썬 패키지 formulas·openpyxl (pip install formulas openpyxl)
시간: 파일 로딩 약 90초 + 판마다 약 20초 (16판 → 약 7분)

동작
1) 웹앱 엔진으로 규칙마다 테스트 판을 만든다: 진행 중, 체크, 체크메이트, P1 승, P2 승,
   그리고 규칙 차이가 드러나는 OXOXOXOXOXO (기본 규칙은 P1 승, 밸런스 패치는 P2 승).
2) 같은 판을 엑셀 수식으로 계산해 상태 코드(계산!B26), 승리 사유(E20), 위험한 수(B24·B25)를 대조한다.
"""
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import formulas
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'OX chess_v3.xlsx'
if not SRC.is_absolute():
    SRC = ROOT / SRC
RULE_NAME = {'classic': '기본 규칙', 'balance': '밸런스 패치'}

# 1) 테스트 판 (웹앱 엔진)
GEN = r"""
const OX = require(process.argv[1]);
const out = {};
for (const preset of ['classic', 'balance']) {
  const r = OX.makeRules(preset); const cases = []; const need = { 1: 2, 2: 2, 3: 1, 4: 2 }; const got = { 1: 0, 2: 0, 3: 0, 4: 0 };
  let seed = 1; const rnd = () => { seed = (seed * 16807) % 2147483647; return seed / 2147483647; };
  while (Object.keys(need).some(k => got[k] < need[k])) {
    const g = new OX.Game(r); const stop = rnd() * 30 | 0;
    while (g.res === null && g.n < stop) g.play(rnd() < .5 ? 0 : 1);
    const st = g.status(); const k = { 1: 1, 2: 2, 3: 3, 4: 4, 5: 4 }[st.code];
    if (got[k] < need[k] && !cases.find(c => c.moves === g.str)) { got[k]++; cases.push({ moves: g.str, code: st.code, winPat: g.winPat, res: g.res, danger: st.danger || [] }); }
  }
  const s = 'OXOXOXOXOXO'; const g = new OX.Game(r); for (const c of s) g.play(c === 'O' ? 0 : 1);
  cases.push({ moves: s, code: g.status().code, winPat: g.winPat, res: g.res, danger: [] });
  out[preset] = cases;
}
console.log(JSON.stringify(out));
"""
cases = json.loads(subprocess.check_output(['node', '-e', GEN, str(ROOT / 'app' / 'engine.js')], text=True))

# 2) 엑셀 수식 계산. 빈 입력 칸은 계산기 입력이 되지 않으므로 공백 한 칸을 넣은 복사본을 쓴다(TRIM 때문에 빈칸과 같음)
wb = openpyxl.load_workbook(SRC)
for col in range(2, 42):
    wb['OX체스'].cell(10, col).value = ' '
probe = Path(tempfile.gettempdir()) / 'oxchess_probe.xlsx'
wb.save(probe)
t = time.time()
xl = formulas.ExcelModel().loads(str(probe)).finish()
print(f'{SRC.name} 로딩 {time.time() - t:.0f}s', flush=True)
name = probe.name   # 계산기 안의 셀 이름에는 파일 이름이 대소문자 그대로 들어간다
cols = [openpyxl.utils.get_column_letter(c) for c in range(2, 42)]

fails = 0
for preset, items in cases.items():
    for case in items:
        m = case['moves']
        inputs = {f"'[{name}]OX체스'!{col}10": (m[i] if i < len(m) else ' ') for i, col in enumerate(cols)}
        inputs[f"'[{name}]OX체스'!W6"] = RULE_NAME[preset]
        idx = {k.upper(): v for k, v in xl.calculate(inputs=inputs).items()}

        def get(sheet, cell):
            v = idx.get(f"'[{name}]{sheet}'!{cell}".upper())
            return v.value[0, 0] if hasattr(v, 'value') else v

        code, reason = int(get('계산', 'B26')), str(get('계산', 'E20'))
        o_bad, x_bad, switch = int(get('계산', 'B24')), int(get('계산', 'B25')), int(get('계산', 'B6'))
        ok = code == case['code'] and switch == (preset == 'balance')
        if case['code'] == 4:
            ok = ok and reason.startswith('P1' if case['res'] == 1 else 'P2') and case['winPat'] in reason
        if case['code'] in (2, 3):
            ok = ok and o_bad == ('O' in case['danger']) and x_bad == ('X' in case['danger'])
        fails += not ok
        print(f"{'✓' if ok else '✗'} {RULE_NAME[preset]:6} {m or '(빈 판)':26} 엑셀 {code} / 엔진 {case['code']} "
              f"위험 O={o_bad} X={x_bad} | {reason or get('OX체스', 'B4')}", flush=True)
os.remove(probe)
print('불일치', fails, '건')
sys.exit(1 if fails else 0)
