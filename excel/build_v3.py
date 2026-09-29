"""OX chess_v2.xlsx → OX chess_v3.xlsx
- 규칙 선택(기본 규칙 / 밸런스 패치) 드롭다운
- 규칙 시트 갱신, 공략 시트 추가
v2의 한 줄 입력 방식과 기존 수식 구조는 그대로 둔다.

기록용: v3를 어떻게 만들었는지 남겨 둔 스크립트. 다음 버전은 v3를 복사해서 고치고 verify_xlsx.py로 검증한다.
사용: python excel/build_v3.py  (저장소 루트의 OX chess_v2.xlsx → OX chess_v3.xlsx)"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'OX chess_v2.xlsx'
DST = ROOT / 'OX chess_v3.xlsx'
WEB = 'https://jeseong-jeong.github.io/oxchess/'

wb = openpyxl.load_workbook(SRC)
ws = wb['OX체스']
calc = wb['계산']
rules = wb['규칙']


def clone_style(dst, src):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


# ---------- OX체스 시트: 규칙 선택 ----------
ws.merge_cells('T6:V6')
ws['T6'] = '규칙'
clone_style(ws['T6'], ws['N6'])                       # '놓은 수' 라벨과 같은 모양
ws.merge_cells('W6:Z6')
ws['W6'] = '기본 규칙'
thin = Side(style='thin', color='FFBFBFBF')
for col in 'WXYZ':
    c = ws[f'{col}6']
    c.fill = PatternFill('solid', fgColor='FFFFF2CC')  # 고를 수 있는 칸 = 연노랑
    c.border = Border(top=thin, bottom=thin, left=thin if col == 'W' else None, right=thin if col == 'Z' else None)
ws['W6'].font = Font(name='Arial', size=11, bold=True)
ws['W6'].alignment = Alignment(horizontal='center', vertical='center')
dv = DataValidation(type='list', formula1='"기본 규칙,밸런스 패치"', allow_blank=False,
                    showErrorMessage=True, errorTitle='규칙', error='기본 규칙 또는 밸런스 패치 중에서 골라 주세요.',
                    promptTitle='규칙 선택', prompt='밸런스 패치: OXO, XOX도 3번이면 P2 승리', showInputMessage=True)
ws.add_data_validation(dv)
dv.add('W6')

ws['A2'] = ("사용법: 아래 한 줄에 1번 칸부터 오른쪽으로 O 또는 X를 한 칸씩 입력하세요. 파랑 칸은 P1, 초록 칸은 P2가 두는 칸이고 "
            "노란 칸이 다음 입력 칸입니다. 오른쪽 위 '규칙' 칸에서 기본 규칙 / 밸런스 패치를 고를 수 있어요. (규칙·공략은 각 시트 참고)")

# ---------- 계산 시트: 규칙 스위치 ----------
calc['A1'] = 'OX체스 계산 시트 (수정하지 마세요. 파란 글씨 3칸만 기준 횟수입니다. 규칙 선택은 OX체스 시트 W6)'
calc['A6'] = '규칙 선택 (1 = 밸런스 패치)'
calc['B6'] = '=IF(OX체스!W6="밸런스 패치",1,0)'
calc['C6'] = 'OX체스 시트 W6에서 선택. 밸런스 패치면 OXO·XOX에도 하우스룰 기준(B4)을 적용'
clone_style(calc['A6'], calc['A5'])
clone_style(calc['C6'], calc['C5'])
calc['B6'].font = copy(calc['B5'].font)
calc['B6'].font = Font(name=calc['B5'].font.name, size=calc['B5'].font.sz, bold=calc['B5'].font.b, color='FF000000')  # 수식 = 검정
calc['B6'].alignment = copy(calc['B5'].alignment)

# 3글자 패턴 행: 32~39 현재, 72~79 O를 둔 뒤, 112~119 X를 둔 뒤 (순서 OOO,OOX,OXO,OXX,XOO,XOX,XXO,XXX)
changed = []
for base in (32, 72, 112):
    for off in range(8):
        r = base + off
        pat = calc.cell(r, 2).value
        if pat in ('OXO', 'XOX'):
            assert calc.cell(r, 4).value == '=$B$3', (r, calc.cell(r, 4).value)
            calc.cell(r, 4).value = '=IF($B$6=1,$B$4,$B$3)'
            changed.append(f'D{r}')
        elif pat in ('OOO', 'XXX'):
            assert calc.cell(r, 4).value == '=$B$4'

# 승리 사유의 '하우스룰' 표시
for r in (20, 21, 22):
    old = calc.cell(r, 5).value
    tag = f'IF(OR(F{r}="OOO",F{r}="XXX"),"·하우스룰","")'
    assert tag in old, old
    calc.cell(r, 5).value = old.replace(
        tag, f'IF(OR(F{r}="OOO",F{r}="XXX",AND($B$6=1,OR(F{r}="OXO",F{r}="XOX"))),"·하우스룰","")')

# ---------- 규칙 시트 다시 쓰기 (스타일은 기존 셀에서 복사) ----------
title_style, head_style, body_style = copy(rules['A1']._style), copy(rules['A3']._style), copy(rules['A4']._style)
for row in rules.iter_rows():
    for c in row:
        c.value = None
lines = [
    ('t', 'OX체스 규칙'),
    None,
    ('h', '■ 진행'),
    ('b', 'P1(선공)과 P2가 번갈아 한 칸씩 O 또는 X를 놓습니다. 매번 O와 X 중 원하는 것을 고릅니다. 놓은 문자는 한 줄로 이어집니다(최대 40칸).'),
    None,
    ('h', '■ 승리 조건'),
    ('b', 'P1: 같은 5글자 패턴이 2번 나열되면 승리 (사이에 다른 문자가 있어도 됨)'),
    ('b', 'P2: 같은 3글자 패턴이 4번 나열되면 승리 (사이에 다른 문자가 있어도 됨)'),
    ('b', '하우스룰: OOO 또는 XXX는 3번만 나열돼도 P2 승리 (예: OOO X OOO XX OOO)'),
    ('b', "밸런스 패치(선택): OXO, XOX도 3번만 나열되면 P2 승리. OX체스 시트 오른쪽 위 '규칙' 칸에서 고릅니다."),
    ('b', '두 조건이 동시에 만족되면 P2 승리'),
    ('b', '패턴은 겹치지 않게 왼쪽부터 셉니다. (예: OOOOOO = OOO 2번, OOOO = OOO 1번)'),
    None,
    ('h', '■ 체크 / 체크메이트'),
    ('b', '체크: 지금 차례인 사람이 O 또는 X 중 한 쪽을 두면 상대가 이기는 상태 (위험한 문자와 이유가 표시됨)'),
    ('b', '체크메이트: O를 두어도 X를 두어도 상대가 이기는 상태'),
    None,
    ('h', '■ 사용법'),
    ('b', "1) OX체스 시트의 한 줄 칸에 1번부터 오른쪽으로 순서대로 O 또는 X를 입력합니다. 빈칸을 건너뛰면 '입력 오류'가 뜹니다."),
    ('b', '2) 상태 패널에서 체크, 체크메이트, 종료 여부를 확인합니다. 게임이 종료되면 입력을 멈추세요.'),
    ('b', '3) 다시 시작하려면 B10:AO10 칸의 내용을 모두 지웁니다.'),
    ('b', "4) 규칙은 OX체스 시트 W6 칸(연노랑)에서 '기본 규칙' / '밸런스 패치' 중에 고릅니다."),
    ('b', '5) 계산 시트는 자동 계산용이므로 수정하지 마세요. 기준 횟수(4, 3, 2)만 계산 시트 B3:B5에서 바꿀 수 있습니다.'),
    None,
    ('h', '■ 웹 버전 (AI 대전 · 같이 하기 · 온라인 대전)'),
    ('link', WEB),
    ('b', "휴대폰에서 열고 '홈 화면에 추가'하면 앱처럼 쓸 수 있어요. 필승 전략은 '공략' 시트를 참고하세요."),
]
for i, item in enumerate(lines, start=1):
    if item is None:
        continue
    kind, text = item
    c = rules.cell(i, 1, text)
    c._style = copy({'t': title_style, 'h': head_style}.get(kind, body_style))
    if kind == 'link':
        c.hyperlink = text
        c.font = Font(name=rules['A4'].font.name, size=rules['A4'].font.sz, color='FF0563C1', underline='single')

# ---------- 공략 시트 ----------
tips = wb.create_sheet('공략', index=wb.sheetnames.index('규칙') + 1)
tips.column_dimensions['A'].width = rules.column_dimensions['A'].width
tips.sheet_view.showGridLines = rules.sheet_view.showGridLines
code_font = Font(name='Consolas', size=11)
T = [
    ('t', 'OX체스 공략'),
    ('b', '모든 수를 끝까지 계산한 결과와 AI끼리 수백 판씩 둔 시뮬레이션에서 나온 내용입니다.'),
    None,
    ('h', '■ 밸런스 요약'),
    ('b', '기본 규칙: 양쪽이 완벽하게 두면 P1 필승 (20수). 둘 다 아무렇게나 두면 P1 65% : P2 35%'),
    ('b', '밸런스 패치: 양쪽이 완벽하게 두면 P2 필승 (22수). 둘 다 아무렇게나 두면 P1 47% : P2 53% (더 공정)'),
    ('b', '선공/후공을 바꿔도 5글자 쪽이 유리합니다. 유불리는 차례가 아니라 승리 조건에서 나옵니다.'),
    ('b', '승부는 초반에 갈립니다. 진 쪽이 처음 실수한 수는 대부분 7~9번째 수였습니다.'),
    None,
    ('h', '■ 공통 원리'),
    ('b', '안전 규칙 먼저: 두면 바로 이기는 수는 두고, 두면 바로 지는 수(체크)는 피합니다.'),
    ('b', '내 목표 패턴의 종류는 줄이고, 상대 목표 패턴의 종류는 늘리세요. 종류가 적을수록 같은 패턴이 빨리 반복됩니다.'),
    None,
    ('h', '■ P1 전략 (기본 규칙 · 필승 쪽)'),
    ('b', "핵심 규칙 '반대로 두기': P2가 방금 둔 글자의 반대를 둡니다."),
    ('b', '  → 판이 OX / XO 짝으로 채워져 OOO·XXX가 절대 생기지 않고(하우스룰 봉쇄), 5글자 패턴 종류가 절반(16가지)으로 줄어 빨리 반복됩니다.'),
    ('b', '첫 수는 O. (X로 시작했다면 아래 표의 O와 X를 모두 뒤집으세요) 9수까지 예외는 이 6자리뿐입니다:'),
    ('c', '   지금까지 판      다음 수       (규칙대로 두면)'),
    ('c', '   OOXXOO        7번째 → O     (X면 짐)'),
    ('c', '   OXOXOX        7번째 → X     (O면 짐)'),
    ('c', '   OOXOXOXO      9번째 → O     (X면 짐)'),
    ('c', '   OOXXOOOX      9번째 → X     (O면 짐)'),
    ('c', '   OXOOXOXO      9번째 → O     (X면 짐)'),
    ('c', '   OXOOXXOO      9번째 → O     (X면 짐)'),
    ('b', '중반부터는 안전 규칙 + 3~4수 앞 읽기. 완벽한 AI 상대 승률: 규칙만 33% → 예외 6개 외우기 64% → 13수까지 예외 15개 85%'),
    ('b', "주의: '반대로 두기'만 고집하면 P2가 X만 계속 둬서 OXOXOX로 끌고 올 때 100% 집니다. 7번째 수 예외 두 개는 꼭 외우세요."),
    None,
    ('h', '■ P2 전략 (기본 규칙 · 이론상 지는 쪽)'),
    ('b', 'P1이 완벽하면 이길 방법이 없습니다. 대신 사람 상대로는 함정이 잘 통합니다.'),
    ('b', "함정 1 'X만 두기'(P1이 O로 시작했을 때): 반대로 두는 P1은 OXOXOX가 되고, 7번째 수에서 자연스러운 O를 두면 집니다."),
    ('b', "함정 2 'P1 따라 하기': OOXXOO로 끌고 가면, 7번째 수에서 자연스러운 X를 두면 집니다."),
    ('b', "함정이 안 통하면 '한 글자만 고집'(O만 계속) + 안전 규칙: P1이 O를 둘 때마다 OOO가 생겨 하우스룰에 가까워집니다."),
    None,
    ('h', '■ P2 전략 (밸런스 패치 · 필승 쪽)'),
    ('b', "핵심 규칙 '세 칸 전 글자 따라 하기': 세 칸 전에 있던 글자를 그대로 둡니다. 판이 3칸 주기로 반복되어 같은 3글자 패턴이 쌓입니다."),
    ('b', 'P1이 O로 시작한 경우 10수까지 예외 6개:'),
    ('c', '   지금까지 판      다음 수'),
    ('c', '   O             2번째 → O'),
    ('c', '   OOO           4번째 → X'),
    ('c', '   OOOXXOO       8번째 → O'),
    ('c', '   OOOXOOOOX     10번째 → X'),
    ('c', '   OOOXXOXXX     10번째 → O'),
    ('c', '   OOXOOXXOX     10번째 → O'),
    ('b', '예외 6개를 외우고 이후 3~4수 앞을 읽으면 완벽한 AI 상대로도 90% 이깁니다.'),
    None,
    ('b', '출처: 웹 버전 저장소의 analysis 폴더 (완전 탐색 솔버, 규칙별 300~600판 시뮬레이션). ' + WEB),
]
for i, item in enumerate(T, start=1):
    if item is None:
        continue
    kind, text = item
    c = tips.cell(i, 1, text)
    c._style = copy({'t': title_style, 'h': head_style}.get(kind, body_style))
    if kind == 'c':
        c.font = code_font

wb.calculation.fullCalcOnLoad = True   # openpyxl은 계산값을 저장하지 않으므로 열 때 다시 계산
wb.active = wb.sheetnames.index('OX체스')
wb.save(DST)
print('saved', DST, '| 바꾼 기준 셀:', ', '.join(changed))
